import openpyxl, sys
try:
    wb = openpyxl.load_workbook('yuntu_pricing.xlsx', data_only=True)
    print('Sheets:', wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f'\n=== {name} === ({ws.max_row} rows x {ws.max_column} cols)')
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(20, ws.max_row), values_only=True)):
            line = ' | '.join([str(c)[:45] if c else '' for c in row][:10])
            print(f'  R{i}: {line}')
except Exception as e:
    print(f'ERROR: {e}', file=sys.stderr)
    import traceback
    traceback.print_exc(file=sys.stderr)
